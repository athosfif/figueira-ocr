"""Read heterogeneous corpora without following links or losing later files."""
from __future__ import annotations
import csv
import io
import os
import stat
from pathlib import Path

SUPPORTED = {'.pdf', '.docx', '.xlsx', '.csv', '.txt', '.log', '.py', '.png', '.jpg', '.jpeg'}


def corpus_root(path):
    path = Path(path).absolute()
    if path.is_symlink() or not path.is_dir():
        raise ValueError('Corpus must be a real directory, not a symbolic link.')
    return path.resolve()


def corpus_files(root):
    for directory, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = sorted(d for d in dirs if not (Path(directory) / d).is_symlink())
        for name in sorted(files):
            path = Path(directory) / name
            if path.is_symlink() or path.suffix.lower() not in SUPPORTED:
                continue
            yield path


def read_regular_file(path, root):
    # O_NOFOLLOW closes the final-component link race. BytesIO parsers never reopen paths.
    path = Path(path)
    if path.is_symlink() or not path.resolve().is_relative_to(root):
        raise ValueError('File escapes the corpus.')
    if not stat.S_ISREG(path.lstat().st_mode):
        raise ValueError('Not a regular corpus file.')
    fd = os.open(path, os.O_RDONLY | getattr(os, 'O_NOFOLLOW', 0) | getattr(os, 'O_NONBLOCK', 0))
    with os.fdopen(fd, 'rb') as stream:
        meta = os.fstat(stream.fileno())
        if not stat.S_ISREG(meta.st_mode) or not meta.st_mode & 0o444:
            raise PermissionError('File is not a readable regular file.')
        cap = int(os.environ.get('RAG_MAX_FILE_BYTES', str(128 * 1024 * 1024)))
        if meta.st_size > cap:
            raise ValueError('File exceeds configured parser safety limit.')
        return stream.read(cap + 1)


def parse_file(path, root, model):
    data = read_regular_file(path, root)
    suffix = path.suffix.lower()
    if suffix in {'.txt', '.log', '.py'}:
        return [('text', data.decode('utf-8-sig', errors='replace'))]
    if suffix == '.csv':
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig', errors='replace'))))
        if not rows:
            return []
        header = rows[0]
        return [(f'row {i}', ' | '.join(f'{header[j] if j < len(header) else j}: {value}'
                                      for j, value in enumerate(row)))
                for i, row in enumerate(rows[1:], 2)]
    if suffix == '.docx':
        # Reading document.xml keeps paragraphs and table cells in original document order.
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read('word/document.xml')
        root_xml = ET.fromstring(xml)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        out = []
        body = root_xml.find('w:body', ns)
        if body is None:
            return out
        def text_of(element):
            return ' '.join(t.text or '' for t in element.findall('.//w:t', ns))
        for i, block in enumerate(body, 1):
            if block.tag.endswith('}p'):
                out.append(text_of(block))
            elif block.tag.endswith('}tbl'):
                rows = [[text_of(cell) for cell in row.findall('w:tc', ns)]
                        for row in block.findall('w:tr', ns)]
                if rows:
                    out.append(' | '.join(rows[0]))
                    for row in rows[1:]:
                        out.append(' | '.join(f'{rows[0][j] if j < len(rows[0]) else j}: {v}'
                                              for j, v in enumerate(row)))
        # Adjacent paragraphs and table context must survive the later overlapping chunk pass.
        return [('document body', '\n'.join(out))]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            out = []
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                for i, row in enumerate(rows, 2):
                    values = [f'{header[j] if j < len(header) and header[j] is not None else j}: {v}'
                              for j, v in enumerate(row) if v is not None]
                    if values:
                        out.append((f'sheet {sheet.title} row {i}', ' | '.join(values)))
                if header:
                    out.insert(0, (f'sheet {sheet.title} header', ' | '.join(str(x) for x in header if x is not None)))
            return out
        finally:
            book.close()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise PermissionError('Encrypted PDF skipped; no password was supplied.')
        out = []
        rendered = None
        try:
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ''
                if not text.strip():
                    import fitz
                    from PIL import Image
                    rendered = rendered or fitz.open(stream=data, filetype='pdf')
                    pix = rendered[i].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    text = model.ocr(Image.frombytes('RGB', (pix.width, pix.height), pix.samples))
                out.append((f'page {i + 1}', text))
            return out
        finally:
            if rendered is not None:
                rendered.close()
    if suffix in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image, ImageOps
        with Image.open(io.BytesIO(data)) as image:
            image = ImageOps.exif_transpose(image).convert('RGB')
            return [('image OCR', model.ocr(image))]
    return []
